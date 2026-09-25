"""Build the report from the reviewed ledger; no network or database access."""

import argparse
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/release-records/releases.json"
OUTPUT = SOURCE.parent / "Conversational_MCQ_Change_and_Deployment_Record.docx"
GENERATOR_VERSION = "change-record-v1"
STATUS = {"verified_live": "已核实上线", "historical_report": "历史核验记录", "pending": "待核实", "failed": "部署未通过"}


def validate(data):
    if data.get("schema_version") != 1:
        raise ValueError("Unsupported ledger schema")
    timestamp = datetime.fromisoformat(data["updated_at"])
    if timestamp.tzinfo is None:
        raise ValueError("updated_at requires a timezone")
    seen = set()
    for release in data["releases"]:
        for field in ("id", "date", "title", "application_commit", "summary", "push_status", "changes", "tests", "data_impact", "limitations", "references"):
            if not release.get(field):
                raise ValueError(f"Missing {field}")
        if release["id"] in seen:
            raise ValueError("Duplicate release ID")
        seen.add(release["id"])
        if not re.fullmatch(r"[0-9a-f]{40}", release["application_commit"]):
            raise ValueError("Full application commit required")
        deployment = release["deployment"]
        if deployment["status"] not in STATUS:
            raise ValueError("Unsupported deployment status")
        if deployment["source_commit"] != release["application_commit"]:
            raise ValueError("Deployment source must match application commit")
        if deployment["status"] == "verified_live":
            for field in ("service", "service_id", "deployment_id", "verified_at", "url", "evidence"):
                if not deployment.get(field):
                    raise ValueError(f"Live deployment missing {field}")
            if datetime.fromisoformat(deployment["verified_at"]).tzinfo is None:
                raise ValueError("Deployment verification needs a timezone")
        for change in release["changes"]:
            if not all(change.get(k) for k in ("area", "problem", "fix", "impact")):
                raise ValueError("Incomplete change entry")
        for test in release["tests"]:
            if not all(test.get(k) for k in ("name", "result", "detail")):
                raise ValueError("Incomplete test evidence")
    if not seen:
        raise ValueError("At least one release required")


def fingerprint(raw):
    return GENERATOR_VERSION + ":" + hashlib.sha256(raw + Path(__file__).read_bytes()).hexdigest()


def set_font(style, size):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + theme), None)
    fonts.set(qn("w:eastAsia"), "Hiragino Sans GB")


def paragraph(doc, text, label=None, style=None):
    p = doc.add_paragraph(style=style)
    if label:
        p.add_run(label + " ").bold = True
    p.add_run(text)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for cell, width in zip(t.rows[0].cells, widths):
        cell.width = Inches(width)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, name in enumerate(headers):
        t.rows[0].cells[i].text = name
    for row in rows:
        cells = t.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = str(value)
    for row_index, row in enumerate(t.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "D9D9D9")
                borders.append(border)
            props.append(borders)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7EEF3")
                props.append(shade)
            for p in cell.paragraphs:
                p.style = doc.styles["Table Text"]
                for run in p.runs:
                    run.bold = row_index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def build(data, identity, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name, size in (("Normal", 10.5), ("Title", 23), ("Subtitle", 12), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        set_font(doc.styles[name], size)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.2
    normal.space_after = Pt(7)
    normal.widow_control = True
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(6)
    for name, size in (("Table Text", 9), ("Reference", 8.5)):
        style = doc.styles.add_style(name, 1)
        style.base_style = doc.styles["Normal"]
        set_font(style, size)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.12
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Conversational MCQ  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "Application changes and deployment evidence"
    doc.core_properties.author = "Conversational MCQ project"
    doc.core_properties.identifier = identity
    stamp = datetime.fromisoformat(data["updated_at"]).astimezone(timezone.utc)
    doc.core_properties.created = doc.core_properties.modified = stamp
    paragraph(doc, data["title"], style="Title")
    paragraph(doc, "持续维护的教学与研究技术档案", style="Subtitle")
    paragraph(doc, "更新日期  " + data["updated_at"] + "  |  时区 America/Edmonton")
    paragraph(doc, data["purpose"])
    doc.add_heading("版本登记", 1)
    table(doc, ["记录", "应用提交", "部署状态"], [
        [r["id"], r["application_commit"][:8], STATUS[r["deployment"]["status"]]] for r in data["releases"]
    ], [2.5, 1.7, 2.7])
    paragraph(doc, "本次详细记录对应 2026 年 9 月 25 日的课堂反馈修复；9 月 22 日内容作为追溯基线。末尾的更早提交索引用于定位开发历史，不代替逐次部署验收。应用提交与随后保存档案的文档提交分开记录。")
    doc.add_heading("阅读和使用边界", 1)
    for text in (
        "测试通过仅支持所覆盖的路径；不代表所有可能的学生行为、模型回答或并发情境都已验证。",
        "观察记录、计算派生值和模型解释必须区分。缺失资料保留为缺失，不为历史记录补造行为。",
        "本档案不保存个人身份、密码或学生原始回答。报告引用应同时注明应用版本和采集或导出版本。",
        "后续修改按项目维护规则追加条目，并重新生成同一 Word 文件；既有条目和限制不因新版本成功而删除。",
    ):
        paragraph(doc, text, style="List Bullet")

    for release in data["releases"]:
        doc.add_heading(release["title"], 1).paragraph_format.page_break_before = True
        paragraph(doc, release["id"] + "  |  " + release["date"])
        paragraph(doc, release["summary"])
        paragraph(doc, release["application_commit"], label="应用提交")
        paragraph(doc, release["push_status"], label="推送状态")
        dep = release["deployment"]
        paragraph(doc, STATUS[dep["status"]] + "  |  " + dep["service"], label="部署状态")
        if dep.get("deployment_id"):
            paragraph(doc, dep["deployment_id"], label="部署编号")
        doc.add_heading("问题与系统性修改", 2)
        for number, change in enumerate(release["changes"], 1):
            doc.add_heading(str(number) + " " + change["area"], 3)
            paragraph(doc, change["problem"], label="问题")
            paragraph(doc, change["fix"], label="修改")
            paragraph(doc, change["impact"], label="影响")
        doc.add_heading("研究数据与计算影响", 2)
        for text in release["data_impact"]:
            paragraph(doc, text, style="List Bullet")
        doc.add_heading("验证结果", 2)
        for test in release["tests"]:
            paragraph(doc, test["result"] + "。" + test["detail"], label=test["name"])
        doc.add_heading("部署证据", 2)
        for text in dep["evidence"]:
            paragraph(doc, text, style="List Bullet")
        if dep.get("url"):
            paragraph(doc, dep["url"], style="Reference")
        doc.add_heading("已知限制", 2)
        for text in release["limitations"]:
            paragraph(doc, text, style="List Bullet")
        doc.add_heading("可追溯文件", 2)
        for path in release["references"]:
            paragraph(doc, path, style="Reference")

    doc.add_heading("相关历史提交索引", 1).paragraph_format.page_break_before = True
    paragraph(doc, "以下按 Git 提交历史整理，仅说明代码变更主题。未逐项重新验证旧版本部署或全部功能；已被后续版本停用的功能不代表当前仍可使用。完整提交可在项目仓库按前缀查询。")
    table(doc, ["日期", "提交前缀", "主题"], data["historical_index"], [1.0, 1.9, 4.0])
    buffer = io.BytesIO()
    doc.save(buffer)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Stable ZIP timestamps avoid binary churn when nothing changed.
    with ZipFile(buffer) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))


def check(output, data, identity):
    doc = Document(output)
    if doc.core_properties.identifier != identity:
        raise ValueError("Word artifact is stale; rebuild from the current ledger and generator")
    text = "\n".join(p.text for p in doc.paragraphs)
    for release in data["releases"]:
        for required in (release["id"], release["application_commit"], release["title"]):
            if required not in text:
                raise ValueError("Missing release content: " + required)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    raw = args.source.read_bytes()
    data = json.loads(raw)
    validate(data)
    identity = fingerprint(raw)
    if not args.check:
        build(data, identity, args.output)
    check(args.output, data, identity)
    print(json.dumps({"output": str(args.output), "releases": len(data["releases"]), "source_identity": identity, "check": "passed"}))


if __name__ == "__main__":
    main()
